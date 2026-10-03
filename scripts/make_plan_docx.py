# -*- coding: utf-8 -*-
"""生成「星牧智控」项目计划书（Word .docx），用于申报上传字段。"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

BASE = r"F:\new大创"
OUT = os.path.join(BASE, "docs", "星牧智控-项目计划书.docx")

doc = Document()

# 全局中文字体
style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

def set_cn(run, name="微软雅黑"):
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), name)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(22); r.bold = True; r.font.color.rgb = RGBColor(0x1F,0x3B,0x57)
    set_cn(r)
    return p

def sub(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x55,0x55,0x55)
    set_cn(r)
    return p

def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.size = Pt(15); r.bold = True; r.font.color.rgb = RGBColor(0x1F,0x3B,0x57)
    set_cn(r)
    return p

def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.font.size = Pt(12.5); r.bold = True
    set_cn(r)
    return p

def body(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    r.font.size = Pt(11)
    set_cn(r)
    return p

def bullet(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.left_indent = Pt(18)
    r = p.add_run("• " + text)
    r.font.size = Pt(11)
    set_cn(r)
    return p

def blank(n=1):
    for _ in range(n):
        doc.add_paragraph()

# ===== 封面信息 =====
blank(2)
title("星牧智控")
sub("——牛只视觉识别、行为分析与健康风险辅助筛查系统")
blank(1)
sub("项目计划书")
blank(1)
for line in [
    "所属领域：农、林、牧、渔业（畜牧业 / 智慧养殖）",
    "项目进展：创意计划阶段",
    "申报组织：大学生创新创业训练计划（创意组）",
]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x44,0x44,0x44)
    set_cn(r)
doc.add_page_break()

# ===== 一、项目概述 =====
h1("一、项目概述")
body("本项目面向规模化牛只养殖场景，研发一套基于机器视觉的牛只识别、行为分析与健康风险辅助筛查系统，用摄像头代替人工巡检，实现牧场全天候、低成本的异常个体发现与健康预警，帮助养殖场在缺人手、缺专业兽医的条件下尽早发现异常牛只。")
body("系统以固定摄像头采集的牛舍图像或视频为输入，先利用深度学习目标检测自动识别画面中的每一头牛，再通过多目标跟踪为每头牛分配持续的身份编号并生成运动轨迹；在轨迹基础上，系统叠加静止、疑似长时间卧倒、离群、活动量下降、区域停留等规则型健康风险判断，并结合群体轨迹异常检测与时序行为识别，输出可复核的结构化告警、个体健康汇总和可视化报告。系统明确设定为辅助筛查用途，只给出风险提示与人工复核建议，不替代兽医的疾病诊断。")

# ===== 二、用户群体与需求 =====
h1("二、用户群体与需求")
bullet("中小型肉牛与奶牛养殖场、规模化牧场：缺人手、缺专业兽医，希望降低巡检与健康管理成本，尽早发现异常个体。")
bullet("畜牧兽医技术服务机构：需要批量、可回溯的结构化风险证据，辅助人工复核与巡检安排。")
bullet("高校畜牧与人工智能相关专业：作为科研与教学的可复现实例，用于可信农业人工智能方向的教学演示。")
body("核心需求是：把牛只健康管理的判断依据从“依赖经验、事后处置”转变为“基于证据、提前预警”，并保证系统结论可解释、可复核、边界清晰。")

# ===== 三、产品与技术方案 =====
h1("三、产品与技术方案")
h2("3.1 系统架构")
body("系统完整链路为：图片或视频输入 → 牛只目标检测 → 多目标跟踪与个体轨迹 → 规则型健康风险 + 群体轨迹异常 + 时序行为识别 → 置信度校准与低可信拒识 → 结构化报告、告警记录与可视化界面。")
h2("3.2 核心技术")
bullet("目标检测：基于 YOLO 系列模型的牛只检测，支持图片与视频输入。")
bullet("多目标跟踪：ByteTrack 与 BoT-SORT，为每头牛分配连续身份编号并生成逐帧轨迹。")
bullet("行为识别：基于 X3D 的时序行为识别，输出中文行为标签、置信度与时间平滑结果。")
bullet("可信校准与拒识：温度校准 + 置信度/熵联合拒识，模型不确定时主动拒绝输出，宁愿“少说”也要“说对”。")
bullet("群体轨迹异常：以同一视频牛群为参照的 Isolation Forest 无监督异常检测，标记活动模式偏离群体的个体。")
bullet("规则型健康风险：静止、疑似长时间卧倒、离群、活动量下降、区域停留等可解释规则。")
h2("3.3 技术栈")
bullet("后端：Python、FastAPI、Uvicorn、Pydantic")
bullet("算法：Ultralytics YOLO、PyTorch、PyTorchVideo、scikit-learn")
bullet("视频与图像：OpenCV、Pillow")
bullet("存储：SQLite、CSV、JSON、HTML、本地媒体文件")
bullet("前端：React、TypeScript、Vite")
bullets_note = body("系统具备模型哈希校验、运行预检、能力降级与结构化实验报告机制，模型缺失或推理失败时明确降级并如实提示，不输出伪造结果。")

# ===== 四、项目愿景与创新点 =====
h1("四、项目愿景与创新点")
body("项目愿景：把牛只健康管理从依赖经验、事后处置，转变为基于证据、提前预警，让可信、可解释、边界清晰的农业人工智能真正落地。")
h2("4.1 创新点一：面向风险的可审计评估链")
body("不只给出一个准确率，而是绑定数据来源与许可、源视频分组、标签映射、清单哈希、模型哈希、训练与测试隔离、预测缓存与失败门控，形成可审计、可复现的评估闭环。")
h2("4.2 创新点二：带校准与拒识的选择性行为识别")
body("系统允许模型在不确定时拒绝输出，用覆盖率、选择性风险与接受样本指标衡量“少做但更可信”的效果，贴近真实牧场辅助筛查需求。")
h2("4.3 创新点三：多层证据融合")
body("同时提供可解释规则、无监督群体异常与学习型行为识别，形成产品与研究基础的多层证据。")
h2("4.4 创新点四：失败可见与能力降级")
body("模型缺失、校准不匹配、推理失败、轨迹不足或产物无效时，如实显示不可用、失败或跳过原因而非伪造完整结果，对农业人工智能的可信部署有现实意义。")

# ===== 五、竞争对手分析 =====
h1("五、竞争对手分析")
body("市面上已有产品多偏重单一功能，如仅做牛只计数、定位或环境温湿度监测，或停留在实验室数据集上、缺少真实场景验证。多数方案存在以下不足：")
bullet("功能单一，缺乏从检测、跟踪到行为与健康告警的完整链路；")
bullet("结论与数据不可追溯，缺少哈希绑定、可复现与失败可见机制；")
bullet("对健康判断过度自信，容易输出缺乏依据的结论；")
bullet("跨场景泛化能力未经真实压力测试。")
body("本项目以完整链路、可审计评估、拒绝过度承诺、承认能力边界的工程严谨性形成差异化优势。")

# ===== 六、实施计划与进度 =====
h1("六、实施计划与进度")
body("项目处于创意计划阶段。当前已完成端到端科研原型的工程闭环（检测、跟踪、行为识别、健康告警、报告与界面的完整链路），并通过自动测试与真实前端构建。")
h2("后续重点推进方向")
bullet("完成跨域迁移适配实验与冻结测试，形成可信的正向跨域结果；")
bullet("补充检测模型的精度报告、推理性能与消融实验；")
bullet("开展真实牧场视频验证，建立个体永久身份与兽医标注真值；")
bullet("完善实时摄像头、边缘部署与人工标注纠错闭环；")
bullet("形成论文与参赛答辩证据链。")

# ===== 七、团队与资源 =====
h1("七、团队与资源")
body("团队具备计算机视觉、深度学习、后端工程与前端开发能力，能够独立完成从算法、工程到界面的完整交付。项目已具备可运行的完整系统，可在本地一键启动演示。")

# ===== 八、预期成果与边界 =====
h1("八、预期成果与边界")
h2("8.1 预期成果")
bullet("一个可运行、可演示的牛只健康辅助筛查系统；")
bullet("一套可审计、可复现的评估协议与实验证据；")
bullet("面向比赛答辩、科研论文与教学演示的完整成果；")
bullet("为真实牧场试点提供接口与方案基础。")
h2("8.2 边界声明")
body("本项目健康结论定位为辅助筛查，用于提示风险与供人工复核，不替代兽医的疾病诊断。跨场景泛化、真实牧场长期数据与医学有效性验证仍在推进中，项目将如实披露能力边界与当前局限。")

blank(1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("（本计划书内容基于项目当前阶段实际情况撰写）")
r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x88,0x88,0x88)
set_cn(r)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("SAVED:", OUT)
print("SIZE:", os.path.getsize(OUT), "bytes")
